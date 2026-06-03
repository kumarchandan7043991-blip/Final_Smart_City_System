from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

slides = [
    {
        'title': 'REAL-TIME TRAFFIC ANALYSIS AND CONGESTION DETECTION USING COMPUTER VISION',
        'subtitle': 'An Intelligent Traffic Perception and Analytics Framework',
        'bullets': [
            'Government Engineering College, Palamu',
            'Department of Computer Science and Engineering',
            'CS702D Major Project',
            'Guide: Mr. Deepak Kr. Ravi',
            'Students: Ashish Kumar [22021440014], Chandan Kumar [22021440016], Chandrashekhar Mahato [22021440017], Ritika Kashish [22021440048]'
        ]
    },
    {
        'title': 'PROBLEM LANDSCAPE',
        'subtitle': 'Urban congestion demands intelligent monitoring',
        'bullets': [
            'Modern cities generate massive traffic movement continuously.',
            'Traditional monitoring systems struggle to interpret dynamic traffic environments.',
            'Increasing vehicle density creates congestion, fuel wastage, pollution, and transportation inefficiency.'
        ]
    },
    {
        'title': 'LIMITATIONS OF EXISTING SYSTEMS',
        'subtitle': 'Observation without intelligence',
        'bullets': [
            'Manual Monitoring: human dependency, delayed response, limited scalability.',
            'Sensor-Based Systems: expensive infrastructure, maintenance complexity, limited flexibility.',
            'Static Signal Systems: fixed timing logic, non-adaptive traffic flow, inefficient congestion handling.'
        ]
    },
    {
        'title': 'CORE ENGINEERING QUESTION',
        'subtitle': 'Transforming raw streams into real intelligence',
        'bullets': [
            'How can raw urban traffic video streams be transformed into real-time actionable traffic intelligence?',
            'The project is not only about vehicle detection.',
            'It is about perception, interpretation, analytics, and intelligent transportation understanding.'
        ]
    },
    {
        'title': 'PROPOSED SYSTEM OVERVIEW',
        'subtitle': 'A complete intelligent perception pipeline',
        'bullets': [
            'Traffic Video Stream → Frame Processing → AI Vehicle Detection → Tracking & Counting → Traffic Analytics → Congestion Intelligence → Visualization & Decision Support.',
            'Transforms raw traffic video into analytical traffic intelligence using Computer Vision and AI.'
        ]
    },
    {
        'title': 'COMPLETE SYSTEM ARCHITECTURE',
        'subtitle': 'Layered intelligent perception pipeline',
        'bullets': [
            'Video Acquisition Layer',
            'Preprocessing Layer',
            'AI Detection Layer',
            'Tracking & Counting Layer',
            'Traffic Analytics Layer',
            'Visualization Layer'
        ]
    },
    {
        'title': 'VIDEO ACQUISITION & PREPROCESSING',
        'subtitle': 'From raw capture to AI-compatible frames',
        'bullets': [
            'Raw traffic video contains computational redundancy and environmental noise.',
            'Preprocessing improves inference speed, computational efficiency, and real-time analytical stability.',
            'Frame Extraction → Noise Reduction → Resolution Optimization → AI-Compatible Frames.'
        ]
    },
    {
        'title': 'AI-BASED VEHICLE DETECTION',
        'subtitle': 'YOLO-enabled real-time perception',
        'bullets': [
            'Input Frame → YOLO Inference Engine → Bounding Box Detection → Vehicle Classification.',
            'YOLO chosen for real-time inference capability, low latency, single-stage detection, and scalable deployment suitability.'
        ]
    },
    {
        'title': 'VEHICLE TRACKING & COUNTING',
        'subtitle': 'Maintaining continuity across video frames',
        'bullets': [
            'Vehicle Detection → Object ID Assignment → Movement Tracking → Temporal Continuity → Accurate Vehicle Counting.',
            'Tracking prevents duplicate counting and preserves continuity across sequential frames.'
        ]
    },
    {
        'title': 'TRAFFIC ANALYTICS ENGINE',
        'subtitle': 'Turning movement into insight',
        'bullets': [
            'Vehicle Data → Density Estimation → Flow Analysis → Congestion Detection → Traffic Intelligence.',
            'Converts raw vehicle movement information into meaningful traffic intelligence.'
        ]
    },
    {
        'title': 'TECHNOLOGY STACK',
        'subtitle': 'Why each technology exists',
        'bullets': [
            'Python: Rapid AI development ecosystem',
            'OpenCV: Efficient computer vision operations',
            'YOLO: Real-time object detection',
            'Flask: Backend communication and integration',
            'MongoDB: Flexible analytical data storage'
        ]
    },
    {
        'title': 'RESULTS & OUTPUTS',
        'subtitle': 'Real-time perception delivering operational outputs',
        'bullets': [
            'Vehicle detection screenshots and tracking outputs demonstrate real-time capability.',
            'Congestion analytics and traffic density outputs show actionable insight.',
            'FPS metrics validate system responsiveness under dynamic traffic conditions.'
        ]
    },
    {
        'title': 'SYSTEM LIMITATIONS',
        'subtitle': 'Engineering risks and operational challenges',
        'bullets': [
            'Low Light → Reduced Visibility',
            'Occlusion → Tracking Complexity',
            'High Resolution Streams → Computational Load'
        ]
    },
    {
        'title': 'ENGINEERING CONTRIBUTIONS',
        'subtitle': 'Moving beyond detection to intelligence',
        'bullets': [
            'The project contribution is not only vehicle detection.',
            'Real-Time Urban Traffic Intelligence Generation through AI-Based Environmental Perception.',
            'Scalable modular architecture, intelligent traffic analytics, AI-driven perception pipeline, congestion interpretation framework.'
        ]
    },
    {
        'title': 'FUTURE SCOPE',
        'subtitle': 'Scaling toward autonomous smart-city systems',
        'bullets': [
            'Real-Time Analytics → Adaptive Traffic Signals → Edge AI Systems → Predictive Traffic Intelligence → Autonomous Smart-City Infrastructure.',
            'The framework can evolve toward adaptive transportation systems and predictive congestion analytics.'
        ]
    },
    {
        'title': 'CONCLUSION',
        'subtitle': 'From static monitoring to real-time intelligence',
        'bullets': [
            'Traditional Monitoring → Static Observation → Limited Intelligence.',
            'AI-Based Systems → Automated Perception → Real-Time Traffic Intelligence.',
            'Establishes a scalable intelligent traffic perception and analytics framework using Computer Vision and AI.'
        ]
    },
    {
        'title': 'THANK YOU',
        'subtitle': 'Questions & Discussion',
        'bullets': [
            'Intelligent transportation infrastructure demo',
            'AI-powered smart-city analytics system',
            'Futuristic engineering architecture showcase',
            'Startup-level technology pitch'
        ]
    }
]

prs = Presentation()
blank_layout = prs.slide_layouts[6]

for slide_data in slides:
    slide = prs.slides.add_slide(blank_layout)
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(3, 7, 14)

    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(1.35))
    title_tf = title_box.text_frame
    title_tf.text = slide_data['title']
    title_tf.paragraphs[0].font.size = Pt(32)
    title_tf.paragraphs[0].font.bold = True
    title_tf.paragraphs[0].font.color.rgb = RGBColor(226, 246, 255)

    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.9), Inches(9), Inches(0.8))
    subtitle_tf = subtitle_box.text_frame
    subtitle_tf.text = slide_data['subtitle']
    subtitle_tf.paragraphs[0].font.size = Pt(18)
    subtitle_tf.paragraphs[0].font.color.rgb = RGBColor(81, 203, 255)

    body_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.7), Inches(9), Inches(4.8))
    body_tf = body_box.text_frame
    body_tf.margin_bottom = Pt(0)
    body_tf.margin_top = Pt(0)
    body_tf.margin_left = Pt(0)
    body_tf.margin_right = Pt(0)
    body_tf.word_wrap = True

    first = True
    for bullet in slide_data['bullets']:
        if first:
            p = body_tf.paragraphs[0]
            p.text = bullet
            first = False
        else:
            p = body_tf.add_paragraph()
            p.text = bullet
            p.level = 0
        p.font.size = Pt(16)
        p.font.color.rgb = RGBColor(196, 216, 232)
        p.font.bold = False
        p.font.name = 'Calibri'
        p.line_spacing = 1.35

    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(7.2), Inches(9), Inches(0.28))
    accent.fill.solid()
    accent.fill.fore_color.rgb = RGBColor(47, 228, 255)
    accent.line.color.rgb = RGBColor(47, 228, 255)
    accent.line.width = Pt(0.5)

prs.save('presentation.pptx')
print('presentation.pptx created successfully.')
